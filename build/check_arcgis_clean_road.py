"""build/check_arcgis_clean_road.py — the Clean Road files on the Reports vs
ArcGIS matrix (v0.50.0).

Drives the SHIPPED path on synthetic data (no corpus files):

  * the Clean Road Intersection and Clean Road Ramp builds from a synthetic
    layer library through their real `consolidate()` — every column the site
    sources comes off the layer the site reads it from, in the site's own print
    forms; the columns the site leaves blank stay blank (the Ramp build's
    odometer is the one exception, written from SHS Ramp Pt.ODMeasure);
  * each build against a site export of the same records is a clean match, a
    changed value counts, a context column never does, and the role gates refuse
    a swapped pair;
  * the Clean Road Highway comparison of the CA HIGHWAYS build against the
    site's export reconciles the spellings that differ (county 'SD.', '02'
    lanes, MM/DD/YYYY dates), keeps the build's unplaceable-span anchor
    non-asserting, and keys a site row with no postmile on its odometer;
  * the registry lists all three Clean Road reports as buildable and
    comparable.

Run with the build venv:
    build\\.venv\\Scripts\\python.exe build\\check_arcgis_clean_road.py
"""
import datetime
import sys
from pathlib import Path

from _checklib import Checker, scripts_path, temp_dir

scripts_path()

from openpyxl import Workbook, load_workbook  # noqa: E402

import arcgis_clean_road_build as acb  # noqa: E402
import arcgis_report_clean_intersection as acri  # noqa: E402
import arcgis_report_clean_ramp as acrr  # noqa: E402
import arcgis_report_intersection_detail as ari  # noqa: E402
import arcgis_reports  # noqa: E402
import clean_highway_columns as chc  # noqa: E402
import clean_road_columns as crc  # noqa: E402
import clean_road_layers as crl  # noqa: E402
import compare_clean_highway_arcgis as cmp_hwy  # noqa: E402
import compare_clean_intersection_arcgis as cmp_inx  # noqa: E402
import compare_clean_ramp_arcgis as cmp_ram  # noqa: E402
import consolidate_clean_road_highway as con_hwy  # noqa: E402
import consolidate_clean_road_intersection as con_inx  # noqa: E402
import consolidate_clean_road_ramp as con_ram  # noqa: E402
from consolidate_clean_highway import UNAVAILABLE_TOKEN  # noqa: E402
from events import Events  # noqa: E402

c = Checker()
D = datetime.datetime
ASOF = "2026-10-05"


class _Quiet(Events):
    def on_log(self, msg):
        pass


def _counts(result):
    co = getattr(result, "comparison_outcome", None)
    k = co.counts if co else None
    return {f: getattr(k, f, None) for f in ("paired_rows", "side_a_only_rows",
                                             "side_b_only_rows", "differing_cells")}


def _q():
    return dict(events=_Quiet(), confirm_overwrite=lambda p: True, mode="preview")


def _sheet(path, title, header, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(list(header))
    for r in rows:
        ws.append([r.get(col) for col in header])
    wb.save(path)


def _span(begin, end, attr, value, county="Orange", route="001", ar=None):
    """A span layer row. `ar` = its (from, to) AR measures, the postmiles by
    default — the Ramp build looks a ramp up by AR, the Intersection build by
    postmile."""
    frm, to = ar or (begin, end)
    return {"District": "District 12", "RouteNum": route, "RouteSuffix": ".",
            "Alignment": "Right", "BeginCounty": county, "EndCounty": county,
            "BeginPMPrefix": ".", "EndPMPrefix": ".", "BeginPMMeasure": begin,
            "EndPMMeasure": end, "LRSFromDate": D(2020, 1, 1), "LRSToDate": None,
            "RouteID": f"SHS_{route}._P", "FromARMeasure": frm, "ToARMeasure": to,
            "InventoryItemStartDate": D(1964, 1, 1), "InventoryItemEndDate": None,
            "Network_Start_Date": D(1964, 1, 1), "Network_End_Date": None,
            attr: value}


_AR_COLS = ["FromARMeasure", "ToARMeasure", "InventoryItemStartDate",
            "InventoryItemEndDate"]


def _library(root, layers):
    """`layers` = [(layer name, columns, rows)] + the INDEX manifest."""
    root.mkdir(parents=True, exist_ok=True)
    made = []
    for n, (layer, cols, rows) in enumerate(layers, start=1):
        name = f"{n:02d}_{layer}.xlsx"
        _sheet(root / name, layer, cols, rows)
        made.append((name, layer, len(rows), len(cols)))
    _sheet(root / crl.INDEX_NAME, "INDEX", list(crl.INDEX_HEADER),
           [dict(zip(crl.INDEX_HEADER, (f, layer, n, k, str(root), "synthetic")))
            for f, layer, n, k in made])


def _overlays(net_route_spans, city_ar=None):
    hg = ari.OVERLAY_LAYERS["HG"]
    city = ari.OVERLAY_LAYERS["CITY"]
    pop = ari.OVERLAY_LAYERS["POP"]
    return [
        (hg[0], ari._SPAN_ID + _AR_COLS + [hg[1]], [_span(0, 50, hg[1], "D- Divided")]),
        (city[0], ari._SPAN_ID + _AR_COLS + [city[1]],
         [_span(0, 1, city[1], "Laguna Beach", ar=city_ar)]),
        (pop[0], ari._SPAN_ID + _AR_COLS + [pop[1]], [_span(0, 50, pop[1], "U- Urban")]),
        ("SHS Inv Network Date", ari._SPAN_ID + ["FromARMeasure", "ToARMeasure",
                                                 "Network_Start_Date", "Network_End_Date",
                                                 "SegOrderId"], net_route_spans),
    ]


def _read_build(path, spec):
    wb = load_workbook(path, read_only=True)
    it = wb[spec.sheet].iter_rows(values_only=True)
    header = [str(x) for x in next(it)]
    rows = [dict(zip(header, r)) for r in it]
    wb.close()
    return rows


def _site_export(folder, spec, key, rows):
    """Per-route site files in the export's own layout, then the app's
    consolidation of them — exactly what the matrix compares the build with."""
    run = folder / spec.key
    run.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = spec.sheet
    ws.append(list(spec.header))
    for r in rows:
        ws.append([r.get(h) or None for h in spec.header])
    wb.save(run / f"2026-10-22 ssor-prod dev-site {spec.key}_route_{key}.xlsx")
    mod = {"clean_highway": con_hwy, "clean_intersection": con_inx,
           "clean_ramp": con_ram}[spec.key]
    out = folder / f"{spec.key}_consolidated 2026-10-22 ssor-prod dev-site.xlsx"
    res = mod.consolidate(events=_Quiet(), confirm_overwrite=lambda p: True,
                          input_dir=run, out_path=out)
    c.check(f"the {spec.label} site export consolidates", res.status == "ok", res.message)
    return out


# --------------------------------------------------------------------------- #
def _inx_row(iid, pm, **over):
    r = {"INTERSECTION_ID": iid, "LRS_DATE_RETIRE": None,
         "InventoryItemStartDate": D(2023, 1, 1),
         "County_Code": "San Bernadino", "District_Code": "District 8",
         "Main_RouteNum": "001", "Main_RouteSuffix": ".", "Main_PMPrefix": ".",
         "Main_PMSuffix": ".", "Main_PMMeasure": pm, "Main_Begin_Date": D(1973, 10, 19),
         "Cross_RouteNum": None, "Cross_RouteSuffix": ".", "Cross_PMPrefix": ".",
         "Cross_PMSuffix": ".", "Cross_PMMeasure": None, "Cross_Begin_Date": D(1974, 5, 1),
         "Intersection_Name": "MAIN ST   LT", "Intersection_Geometry": "T- Tee",
         "Int_Geometry_Begin_Date": D(1964, 1, 1), "Int_Geometry_End_Date": None,
         "Intersection_Control": "B- Stop Signs on Cross Street Only",
         "Int_Control_Begin_Date": D(1973, 10, 19), "Intersection_Lighted_Ind": "Yes",
         "Int_Lighted_Ind_Begin_Date": D(1973, 10, 19), "Int_Date_Of_Record": D(1973, 10, 19)}
    r.update(over)
    return r


def _seg(iid, approach, route_id, direction="Increasing"):
    return {"INTERSECTION_ID": iid, "APPROACH_ID": approach, "ROUTE_ID": route_id,
            "LEG_TRAVEL_DIR": direction}


def _app(approach, lanes, length, left, right, arm):
    return {"APPROACH_ID": approach, "Number_Thru_Lanes": lanes, "FlowCode": "P- Two-Way",
            "N_Distance": length, "Left_Channel": left, "Right_Channel_Ind": right,
            "Signal_Arm_Ind": arm}


def test_intersection(tmp):
    print("Clean Road Intersection: built from the layers, compared with the site's file:")
    lib = tmp / "inx_layers"
    # Every leg of intersection 1 is a state highway (LEG_TYPE would call them
    # all Major); the site picks by ROUTE_ID instead.
    segs = [_seg(1, "7", "SHS_001._P", "Decreasing"), _seg(1, "3", "SHS_001._S"),
            _seg(1, "5", "SHS_001._P"), _seg(1, "9", "SHS_133._P"),
            _seg(2, "11", "SBD_SBD_ELM AVE_P")]
    apps = [_app("7", 9, 999, "X- Other", "No", "No"),
            _app("3", 8, 888, "X- Other", "No", "No"),
            _app("5", 2, 250, "N- None", "No", "No"),
            _app("9", 4, 150, "C- Curbed", "Yes", "Yes"),
            _app("11", 3, 100, "N- None", "No", "No")]
    inx = [_inx_row(1, 0.204, Cross_RouteNum="133", Cross_PMMeasure=2.5),
           _inx_row(2, 3.1, Intersection_Name="ELM AVE"),
           _inx_row(3, None),                                       # no postmile
           _inx_row(4, 4.0, LRS_DATE_RETIRE=D(2024, 1, 1)),         # retired
           _inx_row(5, 4.5, Int_Geometry_Begin_Date=D(2027, 1, 1)),  # not yet built
           _inx_row(6, None, Main_RouteNum=None)]                   # an IM shell
    net = [_span(0, 2, "SegOrderId", 150, county="San Bernardino"),
           _span(2, 9, "SegOrderId", 151, county="San Bernardino"),
           _span(0, 9, "SegOrderId", 9001, county="San Bernardino", route="133")]
    pop = ari.OVERLAY_LAYERS["POP"]
    hg = ari.OVERLAY_LAYERS["HG"]
    city = ari.OVERLAY_LAYERS["CITY"]
    span_cols = ari._SPAN_ID + _AR_COLS
    pop_spans = [dict(_span(0, 3, pop[1], "U- Urban", county="San Bernardino"),
                      EndCounty="Riverside", EndPMPrefix="R", EndPMMeasure=0.5)]
    _library(lib, [(acri.INX_LAYER, acri._INX_COLS, inx),
                   (acri.SEG_LAYER, acri._SEG_COLS, segs),
                   (acri.APP_LAYER, acri._APP_COLS, apps),
                   (hg[0], span_cols + [hg[1]],
                    [_span(0, 50, hg[1], "D- Divided", county="San Bernardino"),
                     _span(0, 50, hg[1], "D- Divided", county="San Bernardino",
                           route="133")]),
                   (city[0], span_cols + [city[1]],
                    [_span(0, 1, city[1], "Laguna Beach", county="San Bernardino")]),
                   (pop[0], span_cols + [pop[1]], pop_spans),
                   ("SHS Inv Network Date",
                    ari._SPAN_ID + ["FromARMeasure", "ToARMeasure", "Network_Start_Date",
                                    "Network_End_Date", "SegOrderId"], net)])
    out = tmp / acri.FILENAME
    res = acri.consolidate(events=_Quiet(), confirm_overwrite=lambda p: True,
                           out_path=out, asof=ASOF, lib_root=lib)
    c.check("the build runs from the layers", res.status == "ok", res.message)
    rows = _read_build(out, crc.INTERSECTION)
    c.check("the site's universe: no-postmile, retired, not-yet-built and shell rows are "
            "left out; the state x state intersection is listed under BOTH routes",
            [(r["Route"], r["INX_CONNECTION_ID"]) for r in rows]
            == [("001", "1"), ("001", "2"), ("133", "1")],
            [(r["Route"], r["INX_CONNECTION_ID"]) for r in rows])
    a, b, x = rows
    c.check("the site's forms: codes, Y/N, MM/DD/YYYY, natural postmile, begin = end, and "
            "the layer's 'San Bernadino' is SBD",
            (a["INX_DESIGN_CODE"], a["INX_CONTROL_CODE"], a["INX_LIGHTED_IND"],
             a["INX_LIGHTED_BEGIN_DATE"], a["INX_BEGIN_PM_AMT"], a["INX_END_PM_AMT"],
             a["INX_COUNTY_CODE"], a["INX_DISTRICT_CODE"], a["INX_ROUTE_NAME"])
            == ("T", "B", "Y", "10/19/1973", "0.204", "0.204", "SBD", "08", "001"), a)
    c.check("legs by route: the mainline is the Increasing _P leg on route 001 with the "
            "lowest approach id; the cross street is the leg on route 133",
            (a["INX_MAIN_LANES_AMT"], a["INX_MAIN_OVERRIDE_LENGTH_AMT"],
             a["INX_CROSS_LANES_AMT"], a["INX_CROSS_SIGNAL_MAST_ARM_IND"],
             a["INX_CROSS_RIGHT_CHANNEL_CODE"], a["INX_CROSS_OVERRIDE_LENGTH_AMT"])
            == ("2", "250", "4", "Y", "Y", "150"), a)
    c.check("...and with no leg on its route the mainline block is blank, the local "
            "road's leg is the cross street",
            not b["INX_MAIN_LANES_AMT"] and b["INX_CROSS_LANES_AMT"] == "3", b)
    c.check("the crossing route's row: located on route 133, crossing 001, its blocks "
            "and begin dates swapped",
            (x["INX_ROUTE_NAME"], x["INX_BEGIN_PM_AMT"], x["INX_X_ROUTE_NAME"],
             x["INX_X_POSTMILE_AMT"], x["INX_MAIN_LANES_AMT"], x["INX_CROSS_LANES_AMT"],
             x["INX_MAIN_BEGIN_DATE"], x["INX_CROSS_BEGIN_DATE"])
            == ("133", "2.5", "001", "0.204", "4", "2", "05/01/1974", "10/19/1973"), x)
    c.check("seg order on the file's route, and on the crossing route at its own postmile",
            [(r["INX_SEG_ORDER_ID"] or "", r["INX_X_SEG_ORDER_ID"] or "") for r in rows]
            == [("150", "9001"), ("151", ""), ("9001", "150")],
            [(r["INX_SEG_ORDER_ID"], r["INX_X_SEG_ORDER_ID"]) for r in rows])
    c.check("overlays at the postmile; City ends at its span; a span that crosses into "
            "another county still covers its start county",
            a["INX_HIGHWAY_GROUP"] == "D" and a["INX_POPULATION_GROUP"] == "U"
            and b["INX_POPULATION_GROUP"] == "U" and a["INX_CITY_CODE"]
            and not b["INX_CITY_CODE"], (a["INX_CITY_CODE"], b["INX_POPULATION_GROUP"]))
    c.check("the columns the site has no source for stay blank; extract date = as-of",
            all(not r[col] for r in rows for col in crc.INTERSECTION.unsourced)
            and a["INX_EXTRACT_DATE"] == "10/05/2026")
    c.check("the build carries its role marker; it is not a ramp build",
            acri.is_arcgis_report(out) and not acrr.is_arcgis_report(out))

    site_rows = [dict(r) for r in rows]
    for r in site_rows:
        r["INX_COUNTY_CODE"] = "SBD"
        r["INX_EXTRACT_DATE"] = "10/22/2026"
    site = _site_export(tmp / "inx_site_same", crc.INTERSECTION, "001", site_rows)
    got = _counts(cmp_inx.compare(str(out), str(site), str(tmp / "ix1.xlsx"), **_q()))
    c.check("the build against the site's file of the same records: a clean match "
            "(the extract date is context)",
            got == {"paired_rows": 3, "side_a_only_rows": 0, "side_b_only_rows": 0,
                    "differing_cells": 0}, got)
    site_rows[1]["INX_CONTROL_CODE"] = "S"
    site_rows[1]["INX_CONNECTION_ID"] = "99"
    site2 = _site_export(tmp / "inx_site_diff", crc.INTERSECTION, "001", site_rows)
    got = _counts(cmp_inx.compare(str(out), str(site2), str(tmp / "ix2.xlsx"), **_q()))
    c.check("a changed control code counts, and so does the connection id (one id space)",
            got["differing_cells"] == 2 and got["paired_rows"] == 3, got)
    res = cmp_inx.compare(str(site), str(out), str(tmp / "ix3.xlsx"), **_q())
    c.check("swapped sides are refused", res.status == "error", res.status)
    written = cmp_inx.compare(str(out), str(site), str(tmp / "ix4.xlsx"),
                              events=_Quiet(), confirm_overwrite=lambda p: True,
                              mode="values")
    notes = " ".join(str(v) for row in load_workbook(tmp / "ix4.xlsx", read_only=True)
                     ["Notes"].iter_rows(values_only=True) for v in row if v)
    c.check("a real workbook names the build's as-of date and the export's day",
            written.status == "ok" and ASOF in notes and "2026-10-22" in notes, notes[:200])


# --------------------------------------------------------------------------- #
def _pt(name, pm, od, **over):
    r = {"District": "District 12", "RouteNum": "001", "RouteSuffix": ".",
         "Alignment": "Right", "County": "Orange", "PMPrefix": ".", "PMMeasure": pm,
         "PMSuffix": ".", "ODMeasure": od, "Ramp_Name": name,
         "InventoryItemStartDate": D(1976, 2, 25), "InventoryItemEndDate": None,
         "RouteID": "SHS_001._S" if over.get("Alignment") == "Left" else "SHS_001._P",
         "ARMeasure": od, "LRSFromDate": D(2020, 1, 1), "LRSToDate": None}
    r.update(over)
    return r


def _attrs(name, **over):
    r = {"Ramp_Name": name, "Ramp_Description": f"001/NB OFF TO {name}",
         "Ramp_Design": "D- Diamond Type Ramp", "Area4_Ind": "Yes",
         "Ramp_On_Off_Ind": "0- OFF", "InventoryItemStartDate": D(1976, 2, 25),
         "InventoryItemEndDate": None, "LRSFromDate": D(2020, 1, 1), "LRSToDate": None}
    r.update(over)
    return r


def test_ramp(tmp):
    print("Clean Road Ramp: built from the layers, compared with the site's file:")
    lib = tmp / "ram_layers"
    pts = [_pt("001_000200", 0.9, 1.9, Alignment="Left"),
           _pt("001_000100", 0.606, 1.6),
           _pt("001_000100", 0.5, 1.5, LRSToDate=D(2021, 1, 1)),        # history
           _pt("001_000300", 2.0, 3.0, InventoryItemStartDate=D(2027, 1, 1))]  # future
    attrs = [_attrs("001_000100"),
             _attrs("001_000200", Ramp_Design="L- Loop-w/o Left Turn",
                    Ramp_On_Off_Ind="1- ON", Area4_Ind="No")]
    dirs = [{"ROUTE": 1, "FROM_": "S", "TO_": "N"}]
    net = [_span(0, 9, "SegOrderId", 150)]
    _library(lib, [(acrr.PT_LAYER, acrr._PT_COLS, pts),
                   (acrr.RAMP_LAYER, acrr._RAMP_COLS, attrs),
                   (acrr.DIR_LAYER, acrr._DIR_COLS, dirs)]
             + _overlays(net, city_ar=(1.5, 1.7)))
    out = tmp / acrr.FILENAME
    res = acrr.consolidate(events=_Quiet(), confirm_overwrite=lambda p: True,
                           out_path=out, asof=ASOF, lib_root=lib)
    c.check("the build runs from the layers", res.status == "ok", res.message)
    rows = _read_build(out, crc.RAMP)
    c.check("one record per live ramp point, ordered by odometer (history and a "
            "not-yet-started ramp left out)",
            [r["RAM_PM_LOC_AMT"] for r in rows] == ["0.606", "0.9"],
            [r["RAM_PM_LOC_AMT"] for r in rows])
    a, b = rows
    c.check("design code + its domain name (whole, as the site writes it), on/off, "
            "Area 4, direction by alignment",
            (a["RAM_DESIGN_CODE"], a["RAM_DESIGN_DESC"], a["RAM_ON_OFF_CODE"],
             a["RAM_AREA_4_IND"], a["RAM_PRIMARY_DIRECTION_CODE"],
             b["RAM_DESIGN_CODE"], b["RAM_ON_OFF_CODE"], b["RAM_AREA_4_IND"],
             b["RAM_PRIMARY_DIRECTION_CODE"])
            == ("D", "D- Diamond Type Ramp", "OFF", "Y", "N", "L", "ON", "N", "S"), (a, b))
    c.check("Highway Group / City / Population are looked up at the ramp's AR "
            "measure (the City span covers AR 1.5-1.7 only); a left-alignment ramp "
            "falls back from its _S route to the _P one",
            a["RAM_CITY_CODE"] and not b["RAM_CITY_CODE"]
            and b["RAM_HIGHWAY_GROUP"] == "D" and b["RAM_SEG_ORDER_ID"] == "150",
            (a["RAM_CITY_CODE"], b["RAM_CITY_CODE"], b["RAM_HIGHWAY_GROUP"]))
    c.check("location, dates, seg order and overlays in the site's forms",
            (a["RAM_DISTRICT_CODE"], a["RAM_COUNTY_CODE"], a["RAM_ROUTE_NAME"],
             a["RAM_BEGIN_DATE"], a["RAM_SEG_ORDER_ID"], a["RAM_HIGHWAY_GROUP"],
             a["RAM_POP_GROUP"], a["RAM_NETWORK_ID"], a["RAM_EXTRACT_DATE"])
            == ("12", "ORA", "001", "02/25/1976", "150", "D", "U", "2", "10/05/2026"), a)
    c.check("the odometer the requirements ask for is written; the other unsourced "
            "columns stay blank",
            a["RAM_BEGIN_OFFSET_AMT"] == "1.6"
            and all(not r[col] for r in rows for col in crc.RAMP.unsourced
                    if col != "RAM_BEGIN_OFFSET_AMT"), a)

    site_rows = [dict(r) for r in rows]
    for r in site_rows:
        r["RAM_BEGIN_OFFSET_AMT"] = ""          # the site leaves it blank
        r["RAM_EXTRACT_DATE"] = "10/22/2026"
    site = _site_export(tmp / "ram_site_same", crc.RAMP, "001", site_rows)
    got = _counts(cmp_ram.compare(str(out), str(site), str(tmp / "r1.xlsx"), **_q()))
    c.check("the build against the site's file of the same ramps: a clean match "
            "(the blank odometer and the extract date are context)",
            got == {"paired_rows": 2, "side_a_only_rows": 0, "side_b_only_rows": 0,
                    "differing_cells": 0}, got)
    site_rows[0]["RAM_DESIGN_DESC"] = "Diamond Type Ramp"
    site_rows.append(dict(site_rows[1], RAM_PM_LOC_AMT="5.5"))
    site2 = _site_export(tmp / "ram_site_diff", crc.RAMP, "001", site_rows)
    got = _counts(cmp_ram.compare(str(out), str(site2), str(tmp / "r2.xlsx"), **_q()))
    c.check("the design description counts here (one domain on both sides), and a ramp "
            "only the site lists is one-sided",
            got == {"paired_rows": 2, "side_a_only_rows": 0, "side_b_only_rows": 1,
                    "differing_cells": 1}, got)


# --------------------------------------------------------------------------- #
def _arc_build(path, rows, skipped=0, marked=0):
    """A minimal CA HIGHWAYS build workbook: the build's sheet + exact header
    and its marker sheet (the role gate)."""
    wb = Workbook()
    ws = wb.active
    ws.title = chc.ARC_SHEET
    ws.append(list(chc.ARC_HEADER))
    for r in rows:
        ws.append([r.get(h) for h in chc.ARC_HEADER])
    mk = wb.create_sheet(chc.ARC_MARKER_SHEET)
    mk.append(["As-of date", ASOF])
    mk.append(["Skipped source spans", skipped])
    mk.append(["Marked anchor cells", marked])
    wb.save(path)


def _thy(begin, end, **kw):
    base = {"THY_SEG_ORDER_ID": "150", "THY_DISTRICT_CODE": "11", "THY_COUNTY_CODE": "SD",
            "THY_ROUTE_NAME": "005", "THY_BEGIN_PM_AMT": begin, "THY_END_PM_AMT": end,
            "THY_LENGTH_MILES_AMT": f"{float(end) - float(begin):.3f}",
            "THY_LT_LANES_AMT": "2", "THY_LEFT_ROAD_EFF_DATE": "1973-07-13",
            "THY_HIGHWAY_GROUP_CODE": "D", "THY_BEGIN_OFFSET_AMT": "1.0",
            "THY_ADT_AMT": "12000", "THY_EXTRACT_DATE": "2026-10-05"}
    base.update(kw)
    return base


def test_highway(tmp):
    print("Clean Road Highway: the CA HIGHWAYS build vs the site's file:")
    build = tmp / "clean_highway_built.xlsx"
    _arc_build(build, [_thy("0", "0.305"), _thy("0.305", "0.9"),
                       _thy("0.9", "1.2", THY_LT_LANES_AMT=UNAVAILABLE_TOKEN)],
               skipped=1, marked=1)
    site_rows = [_thy("0", "0.305", THY_COUNTY_CODE="SD.", THY_LT_LANES_AMT="02",
                      THY_LEFT_ROAD_EFF_DATE="07/13/1973", THY_BEGIN_OFFSET_AMT="7.0",
                      THY_ADT_AMT="", THY_EXTRACT_DATE="10/22/2026"),
                 _thy("0.305", "0.9", THY_COUNTY_CODE="SD.", THY_LT_LANES_AMT="03",
                      THY_LEFT_ROAD_EFF_DATE="07/13/1973", THY_ADT_AMT=""),
                 _thy("0.9", "1.2", THY_COUNTY_CODE="SD.", THY_LT_LANES_AMT="04",
                      THY_LEFT_ROAD_EFF_DATE="07/13/1973", THY_ADT_AMT=""),
                 _thy("1.2", "1.5", THY_COUNTY_CODE="SD.", THY_BEGIN_PM_AMT="",
                      THY_END_PM_AMT="", THY_BEGIN_OFFSET_AMT="2.2",
                      THY_LEFT_ROAD_EFF_DATE="07/13/1973")]
    site = _site_export(tmp / "hwy_site", crc.HIGHWAY, "005", site_rows)
    out = tmp / "h1.xlsx"
    res = cmp_hwy.compare(str(build), str(site), str(out), events=_Quiet(),
                          confirm_overwrite=lambda p: True, mode="values")
    got = _counts(res)
    c.check("'SD.' = 'SD', '02' = '2', MM/DD/YYYY = ISO pair the segments; the "
            "offset, ADT and extract date are context; the unavailable anchor asserts "
            "nothing; the one real lane change counts; the no-postmile row is one-sided",
            got == {"paired_rows": 3, "side_a_only_rows": 0, "side_b_only_rows": 1,
                    "differing_cells": 1}, got)
    notes = " ".join(str(v) for row in load_workbook(out, read_only=True)["Notes"]
                     .iter_rows(values_only=True) for v in row if v)
    c.check("the Notes state the vintage and name the row with no postmile",
            ASOF in notes and "NO postmile" in notes and "2026-10-22" in notes, notes[:300])
    res = cmp_hwy.compare(str(site), str(build), str(tmp / "h2.xlsx"), **_q())
    c.check("swapped sides are refused", res.status == "error", res.status)


def main():
    print("the registry:")
    eds = {e.key: e for e in arcgis_reports.editions()}
    c.check("all six Clean Road editions compare against their report's build",
            all(eds[k].comparable for k in ("clean_highway", "clean_highway_pdf",
                                            "clean_intersection", "clean_intersection_pdf",
                                            "clean_ramp", "clean_ramp_pdf")))
    c.check("both new builds declare the layers they refuse without",
            {"SHS Ramp Pt", "SHS Ramp", "Route Direction", "SHS Inv Network Date"}
            <= set(acrr.REQUIRED_LAYERS)
            and "SHS Inv Network Date" in acri.REQUIRED_LAYERS
            and set(acri.REQUIRED_LAYERS) <= set(crl.EXPECTED_LAYERS)
            and set(acrr.REQUIRED_LAYERS) <= set(crl.EXPECTED_LAYERS))
    c.check("the builds write the site export's exact header",
            acb.ROUTE_COL == "Route" and acri.SPEC is crc.INTERSECTION
            and acrr.SPEC is crc.RAMP)
    with temp_dir("tsmis_agcr_") as tmp:
        tmp = Path(tmp)
        test_intersection(tmp)
        test_ramp(tmp)
        test_highway(tmp)
    return c.summary()


if __name__ == "__main__":
    sys.exit(main())
