"""Clean Road Highway — the ArcGIS-built CA HIGHWAYS workbook vs the SITE's
Clean Road Highway export (v0.50.0).

The Reports vs ArcGIS matrix's Clean Road Highway rows. Side A is OUR CA
HIGHWAYS table built from the ArcGIS layer library (`consolidate_clean_highway`,
the ArcGIS reports tab); side B is the site's own Clean Road Highway file for one
day, consolidated — the Excel export or the print read back from its PDF, both
the same shape (the PDF-vs-Excel self-check proves it), so one comparator serves
both editions.

This one is TSMIS vs TSMIS: both sides are built from the same LRS layers, so a
difference is a gap in OUR reconstruction or a finding about how the site builds
its file — never migration drift. Values are compared on the vs-TSN comparison's
FORMAT-only reconciliations, because the two sides really do spell values
differently (the site prints MM/DD/YYYY, '02' lanes and 'LA.' for the two-letter
counties; the build writes the TSN extract's shapes): dates as ISO dates,
amounts as canonical numbers, county codes without the trailing period, a blank
toll/forest code as 0, and the equation marker 'E' the site writes into the
postmile-suffix column is no roadbed for the key. Nothing semantic is bridged.

Rows pair exactly as the site-vs-TSN comparison pairs them (route + county + PM
prefix + BEGIN postmile + roadbed), through the same `compare_clean_road_tsn`
projection, so the two comparisons of one export can never key it differently.
A site row with no postmile keys on its odometer instead and is disclosed on the
Notes sheet (`compare_clean_road_tsn.UNPLACED`).

CONTEXT (shown, never counted) is the vs-TSN set: the columns the site has no
source for (blank in its file by design — the build may still carry a value, the
ADT trio for one), TSN's bookkeeping columns, the two offset columns (the site
writes the LRS odometer, the build its PM-continued cumulative),
THY_LAST_SIG_CHG_DATE, THY_FUNCTIONAL_CLASS_CODE and the extract date. The
build's unplaceable-span anchors stay NON-ASSERTING here exactly as they are
against TSN.

VINTAGE is the first thing on the Notes sheet: the build is a reconstruction as
of a date, the export is one day's file; across a gap the counts include network
change. Console-free; engine in compare_core.
"""
import logging
import re
from dataclasses import replace
from pathlib import Path

try:
    from openpyxl import load_workbook
    _DEPS_OK = True
except ImportError:
    _DEPS_OK = False

import clean_highway_columns as chc
import compare_clean_highway_tsn as cht
import compare_clean_road_tsn as crt
import compare_tsn_common as ctc
import consolidate_clean_highway as cch
from consolidate_clean_highway import UNAVAILABLE_TOKEN
from paths import today_str
from pdf_table_lib import pdf_source_marker_state

log = logging.getLogger("tsmis.compare")

REPORT_NAME = "Clean Road Highway — ArcGIS vs TSMIS"
SIDE_A = "ArcGIS"
SIDE_B = "TSMIS"
file_a_label = SIDE_A
file_b_label = SIDE_B

_PROFILE = crt.profile_for("clean_highway")
_SPEC = _PROFILE.spec
_ARC_INDEX = {name: i for i, name in enumerate(chc.ARC_HEADER)}
SHARED_HEADER = list(_SPEC.header)
CONTEXT_FIELDS = tuple(_PROFILE.context)
# What the site's export adds to the build's own context: the columns its file
# leaves blank by design (`clean_road_columns`) and its extract date.
EXPORT_CONTEXT_FIELDS = tuple(c for c in CONTEXT_FIELDS
                              if c not in cch.CONTEXT_COLUMNS)


def _s(v):
    return "" if v is None else str(v).strip()


def _open(path):
    name = Path(path).name
    try:
        return name, load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"Could not open {name}: {type(e).__name__}: {e}")


def _load_build(path, collapse):
    """Side A: OUR CA HIGHWAYS workbook — marker-gated, the exact build header —
    projected onto the site export's 74 columns by name and run through the
    same projection the site side uses. A build-only column has no site
    counterpart and is left out; an unplaceable-span anchor keeps its reserved
    token verbatim so the schema can hold it non-asserting."""
    name, wb = _open(path)
    try:
        if chc.ARC_MARKER_SHEET not in wb.sheetnames:
            raise ValueError(
                f"The {SIDE_A} side is not an ArcGIS-built Clean Road Highway "
                f"workbook (no '{chc.ARC_MARKER_SHEET}' sheet):\n{path}\n\n"
                "Build it on the ArcGIS ▸ ArcGIS reports tab first.")
        if chc.ARC_SHEET not in wb.sheetnames:
            raise ValueError(f"{name} has no '{chc.ARC_SHEET}' sheet — rebuild "
                             "the Clean Road Highway report.")
        it = wb[chc.ARC_SHEET].iter_rows(values_only=True)
        header = [_s(c) for c in (next(it, None) or ())]
        if header != list(chc.ARC_HEADER):
            raise ValueError(
                f"{name} does not carry the exact {len(chc.ARC_HEADER)}-column "
                "build header — rebuild the Clean Road Highway report with this "
                "version.")
        hint = f"{name} ({chc.ARC_SHEET})"
        rows = []
        for r in it:
            if not ctc.row_has_data(r):
                continue
            r = list(r) + [None] * (len(chc.ARC_HEADER) - len(r))
            vals = [r[_ARC_INDEX[c]] for c in _SPEC.header]
            out = crt._row(_PROFILE, vals, "tsn", hint, collapse=collapse)
            for i, raw in enumerate(vals):
                if raw == UNAVAILABLE_TOKEN and _SPEC.header[i] != _SPEC.pm_col:
                    out[1 + i] = UNAVAILABLE_TOKEN
            rows.append(out)
        return rows
    finally:
        wb.close()


def _is_print(path):
    return pdf_source_marker_state(path) != 0


def _load_pair(build_path, export_path):
    """Both sides, with the ROLE of each enforced: side A must be the app's
    build, side B a consolidated site export (`load_site` refuses the build).
    A print-edition export folds whitespace runs on both sides, as the print's
    HTML does."""
    pdf = _is_print(export_path)
    rows_a = _load_build(build_path, collapse=pdf)
    rows_b = crt.load_site(export_path, _PROFILE, pdf=pdf, collapse=pdf)
    return ctc.same_source_render_rows(rows_a), ctc.same_source_render_rows(rows_b), None


# --------------------------------------------------------------------------- #
# Notes
# --------------------------------------------------------------------------- #
def _export_day(path):
    m = re.search(r"(\d{4}-\d{2}-\d{2})", Path(path).name)
    return m.group(1) if m else ""


def _facts(build_path):
    try:
        return cch.report_facts(build_path)
    except Exception as e:   # silent-ok: the Notes sheet is documentation; an unreadable marker costs one stated date, never the comparison (the loaders gate the workbooks themselves)
        log.info("clean-road arcgis notes: build facts unreadable (%s: %s)",
                 type(e).__name__, e)
        return {}


def _notes_lines(build_path, export_path, unplaced):
    facts = _facts(build_path)
    asof = facts.get("asof") or "(unknown)"
    drop = facts.get("drop_exported") or "(unknown)"
    day = _export_day(export_path) or "an unrecorded day"
    edition = "print (read back from its PDF)" if _is_print(export_path) else "Excel export"
    return (
        f"Side A ({SIDE_A}) is OUR CA HIGHWAYS table, rebuilt from the ArcGIS "
        f"layer library as of {asof} (layers exported {drop}). Side B "
        f"({SIDE_B}) is the site's own Clean Road Highway file — its {edition} "
        f"of {day}, consolidated.",
        "VINTAGE: these two dates must match for the counts to measure "
        "correctness. Across a gap they include network change — refresh the "
        "layers (ArcGIS ▸ Layers), rebuild the report as of the export day "
        "(ArcGIS ▸ ArcGIS reports) and compare again.",
        "Both sides are built from the same LRS layers, so a difference here is "
        "a gap in our reconstruction or a finding about how the site builds its "
        "file — never migration drift.",
        "Rows pair on route + county + PM prefix + BEGIN postmile + roadbed (the "
        "R/L/X postmile suffix) — the key the site-vs-TSN comparison uses. The "
        "END postmile and length are compared fields. The site cuts a segment "
        "at every event point of its Highway Log engine; the build cuts one at "
        "every layer breakpoint, so most one-sided rows are the two "
        "segmentations disagreeing, not missing road.",
        "Format-only reconciliations (the two sides spell values differently): "
        "dates compare as ISO dates; amounts as canonical numbers ('02' = '2'); "
        "county codes drop the trailing period the site writes for two-letter "
        "counties ('SD.' = 'SD'); a blank THY_TOLL_FOREST_CODE is code 0; the "
        "equation marker 'E' the site writes into THY_PM_SUFFIX_CODE is no "
        "roadbed for the key (the column itself still compares). A print-edition "
        "export folds runs of spaces on both sides, as its HTML does.",
        "CONTEXT (shown, never counted): the 22 columns the site has no source "
        "for (blank in its file by design — the build may still carry a value, "
        "e.g. the ADT trio), TSN's bookkeeping columns, the two offset columns "
        "(the site writes the LRS odometer, the build its PM-continued "
        "cumulative), THY_LAST_SIG_CHG_DATE, THY_FUNCTIONAL_CLASS_CODE and the "
        "extract date.",
    ) + tuple(unplaced) + (
        "Column match table — each THY column · the ArcGIS layer (source "
        "column) the build paints it from · whether the vs-TSN comparison "
        "counts it (this comparison counts the same columns):",
    ) + cht._provenance_table_lines()


# --------------------------------------------------------------------------- #
_BASE_SCHEMA = crt._schema(
    _PROFILE, SIDE_A, SIDE_B,
    notes_title=f"{REPORT_NAME}: comparison notes", notes=(),
    context=_PROFILE.context, representation=_PROFILE.text)
_BASE_SCHEMA = replace(
    _BASE_SCHEMA, report_name=REPORT_NAME, sides_noun="builds",
    one_sided_note_extra=(" (segments one side cuts at a postmile the other "
                          "doesn't — mostly the two segmentations, or network "
                          "change across the two dates)"))


def _schema_for(build_path, export_path, found):
    """The per-run schema: the Notes sheet reads what the loader found, and a
    build that recorded unplaceable source spans opts into the non-asserting
    anchor token (the vs-TSN comparison's HF-01 rule)."""
    def notes(wb):
        return ctc.make_notes_writer(
            f"{REPORT_NAME}: comparison notes",
            _notes_lines(build_path, export_path, found.get("unplaced", ())))(wb)

    schema = replace(_BASE_SCHEMA, legend_writer=notes)
    skipped, marked, _withheld = cht._build_skip_facts(build_path)
    if skipped or marked:
        schema = replace(schema, unavailable_rule=(
            UNAVAILABLE_TOKEN,
            lambda: "SOURCE COVERAGE — " + cht._coverage_sentence(
                skipped, marked, None)))
    return schema


def suggest_name(_build_path=None):
    return f"ArcGIS_vs_TSMIS_CleanRoadHighway_Comparison {today_str()}.xlsx"


def compare(build_path, export_path, out_path, events=None,
            confirm_overwrite=None, mode="formulas", commit_guard=None):
    """Build the ArcGIS-vs-TSMIS Clean Road Highway discrepancy workbook(s)."""
    found = {}

    def loader(a, b):
        rows_a, rows_b, extra = _load_pair(a, b)
        found["unplaced"] = crt.unplaced_note_lines(
            _SPEC, rows_a, rows_b, SIDE_A, SIDE_B)
        return rows_a, rows_b, extra

    return ctc.run_files_compare(
        _schema_for(build_path, export_path, found), build_path, export_path,
        out_path, banner=f"Clean Road Highway Comparison — {SIDE_A} vs {SIDE_B}",
        has_route=True, loader=loader, deps_ok=_DEPS_OK,
        deps_msg="Required components are missing (openpyxl).",
        side_a=SIDE_A, side_b=SIDE_B,
        events=events, confirm_overwrite=confirm_overwrite, mode=mode,
        commit_guard=commit_guard)
